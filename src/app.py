from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
import joblib
import pandas as pd
import sqlite3
import io
import os
from datetime import datetime

# Excel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

# PDF
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer
)


# =====================================================
# FASTAPI APP
# =====================================================

app = FastAPI(
    title="Student Dropout Risk Prediction API",
    description="AI-based Student Dropout Risk Prediction System",
    version="1.0"
)


# =====================================================
# PATHS
# =====================================================

MODEL_PATH = "models/logistic_regression_model.pkl"
DATABASE_NAME = "student_predictions.db"
BULK_OUTPUT_FILE = "bulk_predictions.csv"


# =====================================================
# LOAD MODEL
# =====================================================

try:
    model = joblib.load(MODEL_PATH)
    print("Model loaded successfully.")
except Exception as e:
    print("Error loading model:", e)
    model = None


# =====================================================
# REQUIRED COLUMNS
# =====================================================

REQUIRED_COLUMNS = [
    "Age",
    "Gender",
    "Year_of_Study",
    "Attendance_Percentage",
    "Average_Grade",
    "Failed_Subjects",
    "Assignments_Completed_Percentage",
    "Engagement_Score",
    "Study_Hours_Per_Day",
    "Online_Learning_Hours_Per_Week",
    "Family_Income_INR",
    "Scholarship",
    "Part_Time_Job",
    "Financial_Stress_Score",
    "Distance_From_College_KM",
    "Internet_Access",
    "Parental_Support",
    "Extracurricular_Activities",
    "Disciplinary_Issues"
]


# =====================================================
# PYDANTIC MODEL
# =====================================================

class StudentData(BaseModel):

    Age: int
    Gender: str
    Year_of_Study: int
    Attendance_Percentage: float
    Average_Grade: float
    Failed_Subjects: int
    Assignments_Completed_Percentage: float
    Engagement_Score: float
    Study_Hours_Per_Day: float
    Online_Learning_Hours_Per_Week: float
    Family_Income_INR: float
    Scholarship: str
    Part_Time_Job: str
    Financial_Stress_Score: float
    Distance_From_College_KM: float
    Internet_Access: str
    Parental_Support: str
    Extracurricular_Activities: str
    Disciplinary_Issues: str


# =====================================================
# HELPER FUNCTION
# =====================================================

def sqlite_value(value):

    if pd.isna(value):
        return None

    if hasattr(value, "item"):
        return value.item()

    return value


# =====================================================
# HOME PAGE
# =====================================================

@app.get("/")
def home():

    return FileResponse(
        "frontend/index.html"
    )


# =====================================================
# DASHBOARD PAGE
# =====================================================

@app.get("/dashboard")
def dashboard():

    return FileResponse(
        "frontend/dashboard.html"
    )


# =====================================================
# PREDICTION HISTORY PAGE
# =====================================================

@app.get("/prediction-history")
def prediction_history():

    return FileResponse(
        "frontend/prediction-history.html"
    )


# =====================================================
# STUDENT DETAILS PAGE
# =====================================================

@app.get("/student-details")
def student_details():

    return FileResponse(
        "frontend/student-details.html"
    )

@app.get("/settings")
def settings():
    return FileResponse("frontend/settings.html")


@app.get("/about")
def about():
    return FileResponse("frontend/about.html")

# =====================================================
# HEALTH CHECK
# =====================================================

@app.get("/api/health")
def health_check():

    return {
        "status": "healthy",
        "model_loaded": model is not None
    }


# =====================================================
# SINGLE STUDENT PREDICTION
# =====================================================

@app.post("/predict")
def predict_student(student: StudentData):

    if model is None:
        raise HTTPException(
            status_code=500,
            detail="Model could not be loaded."
        )

    try:

        input_data = pd.DataFrame(
            [student.model_dump()]
        )

        prediction = model.predict(input_data)[0]

        probability = model.predict_proba(
            input_data
        )[0][1]

        dropout_probability = round(
            float(probability) * 100,
            2
        )

        # Risk level
        if dropout_probability < 30:
            risk_level = "LOW"

        elif dropout_probability < 60:
            risk_level = "MEDIUM"

        else:
            risk_level = "HIGH"

        # Prediction label
        if prediction == 1:

            prediction_label = (
                "At Risk of Dropout"
            )

        else:

            prediction_label = (
                "Not Currently Predicted as Dropout"
            )

        # =================================================
        # SAVE TO DATABASE
        # =================================================

        connection = sqlite3.connect(
            DATABASE_NAME
        )

        cursor = connection.cursor()

        cursor.execute(
            """
            INSERT INTO predictions (
                age,
                gender,
                year_of_study,
                attendance_percentage,
                average_grade,
                failed_subjects,
                assignments_completed_percentage,
                engagement_score,
                study_hours_per_day,
                online_learning_hours_per_week,
                family_income_inr,
                scholarship,
                part_time_job,
                financial_stress_score,
                distance_from_college_km,
                internet_access,
                parental_support,
                extracurricular_activities,
                disciplinary_issues,
                dropout_probability,
                risk_level,
                prediction
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                student.Age,
                student.Gender,
                student.Year_of_Study,
                student.Attendance_Percentage,
                student.Average_Grade,
                student.Failed_Subjects,
                student.Assignments_Completed_Percentage,
                student.Engagement_Score,
                student.Study_Hours_Per_Day,
                student.Online_Learning_Hours_Per_Week,
                student.Family_Income_INR,
                student.Scholarship,
                student.Part_Time_Job,
                student.Financial_Stress_Score,
                student.Distance_From_College_KM,
                student.Internet_Access,
                student.Parental_Support,
                student.Extracurricular_Activities,
                student.Disciplinary_Issues,
                dropout_probability,
                risk_level,
                prediction_label
            )
        )

        connection.commit()

        connection.close()

        return {
            "prediction": prediction_label,
            "dropout_probability": dropout_probability,
            "risk_level": risk_level
        }

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )


# =====================================================
# BULK CSV PREDICTION
# =====================================================

@app.post("/bulk-predict")
async def bulk_predict(
    file: UploadFile = File(...)
):

    if model is None:

        raise HTTPException(
            status_code=500,
            detail="Model could not be loaded."
        )

    try:

        contents = await file.read()

        df = pd.read_csv(
            io.BytesIO(contents)
        )

        # Check required columns
        missing_columns = [
            column
            for column in REQUIRED_COLUMNS
            if column not in df.columns
        ]

        if missing_columns:

            raise HTTPException(
                status_code=400,
                detail={
                    "message": "Missing required columns.",
                    "missing_columns": missing_columns
                }
            )

        input_df = df[
            REQUIRED_COLUMNS
        ].copy()

        # Predictions
        predictions = model.predict(
            input_df
        )

        probabilities = model.predict_proba(
            input_df
        )[:, 1]

        df["Dropout_Probability"] = (
            probabilities * 100
        ).round(2)

        df["Risk_Level"] = df[
            "Dropout_Probability"
        ].apply(
            lambda x:
                "LOW"
                if x < 30
                else
                "MEDIUM"
                if x < 60
                else
                "HIGH"
        )

        df["Prediction"] = [
            "At Risk of Dropout"
            if prediction == 1
            else
            "Not Currently Predicted as Dropout"
            for prediction in predictions
        ]

        # =================================================
        # SAVE TO DATABASE
        # =================================================

        connection = sqlite3.connect(
            DATABASE_NAME
        )

        cursor = connection.cursor()

        for _, row in df.iterrows():

            cursor.execute(
                """
                INSERT INTO predictions (
                    age,
                    gender,
                    year_of_study,
                    attendance_percentage,
                    average_grade,
                    failed_subjects,
                    assignments_completed_percentage,
                    engagement_score,
                    study_hours_per_day,
                    online_learning_hours_per_week,
                    family_income_inr,
                    scholarship,
                    part_time_job,
                    financial_stress_score,
                    distance_from_college_km,
                    internet_access,
                    parental_support,
                    extracurricular_activities,
                    disciplinary_issues,
                    dropout_probability,
                    risk_level,
                    prediction
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    sqlite_value(row["Age"]),
                    sqlite_value(row["Gender"]),
                    sqlite_value(row["Year_of_Study"]),
                    sqlite_value(row["Attendance_Percentage"]),
                    sqlite_value(row["Average_Grade"]),
                    sqlite_value(row["Failed_Subjects"]),
                    sqlite_value(row["Assignments_Completed_Percentage"]),
                    sqlite_value(row["Engagement_Score"]),
                    sqlite_value(row["Study_Hours_Per_Day"]),
                    sqlite_value(row["Online_Learning_Hours_Per_Week"]),
                    sqlite_value(row["Family_Income_INR"]),
                    sqlite_value(row["Scholarship"]),
                    sqlite_value(row["Part_Time_Job"]),
                    sqlite_value(row["Financial_Stress_Score"]),
                    sqlite_value(row["Distance_From_College_KM"]),
                    sqlite_value(row["Internet_Access"]),
                    sqlite_value(row["Parental_Support"]),
                    sqlite_value(row["Extracurricular_Activities"]),
                    sqlite_value(row["Disciplinary_Issues"]),
                    sqlite_value(row["Dropout_Probability"]),
                    sqlite_value(row["Risk_Level"]),
                    sqlite_value(row["Prediction"])
                )
            )

        connection.commit()

        connection.close()

        # Save bulk output
        df.to_csv(
            BULK_OUTPUT_FILE,
            index=False
        )

        return {
            "message": "Bulk prediction completed successfully.",
            "total_students": len(df),
            "results": df.to_dict(
                orient="records"
            )
        }

    except HTTPException:
        raise

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )


# =====================================================
# DOWNLOAD BULK PREDICTIONS
# =====================================================

@app.get("/bulk_predictions.csv")
def download_bulk_predictions():

    if not os.path.exists(
        BULK_OUTPUT_FILE
    ):

        raise HTTPException(
            status_code=404,
            detail="Bulk prediction file not found."
        )

    return FileResponse(
        BULK_OUTPUT_FILE,
        media_type="text/csv",
        filename="bulk_predictions.csv"
    )


# =====================================================
# GET ALL PREDICTIONS
# =====================================================

@app.get("/predictions")
def get_predictions():

    try:

        connection = sqlite3.connect(
            DATABASE_NAME
        )

        connection.row_factory = sqlite3.Row

        cursor = connection.cursor()

        cursor.execute(
            """
            SELECT
                id,
                age,
                gender,
                year_of_study,
                attendance_percentage,
                average_grade,
                failed_subjects,
                assignments_completed_percentage,
                engagement_score,
                study_hours_per_day,
                online_learning_hours_per_week,
                family_income_inr,
                scholarship,
                part_time_job,
                financial_stress_score,
                distance_from_college_km,
                internet_access,
                parental_support,
                extracurricular_activities,
                disciplinary_issues,
                dropout_probability,
                risk_level,
                prediction,
                created_at
            FROM predictions
            ORDER BY id DESC
            """
        )

        rows = cursor.fetchall()

        connection.close()

        return [
            dict(row)
            for row in rows
        ]

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )


# =====================================================
# EXPORT PREDICTION HISTORY
# =====================================================

@app.get("/export-predictions")
def export_predictions(
    risk: str = "ALL",
    format: str = "csv"
):

    # -------------------------------------------------
    # VALIDATE RISK
    # -------------------------------------------------

    risk = risk.upper()
    format = format.lower()

    allowed_risks = [
        "ALL",
        "LOW",
        "MEDIUM",
        "HIGH"
    ]

    allowed_formats = [
        "csv",
        "xlsx",
        "pdf"
    ]

    if risk not in allowed_risks:

        raise HTTPException(
            status_code=400,
            detail="Invalid risk level."
        )

    if format not in allowed_formats:

        raise HTTPException(
            status_code=400,
            detail="Invalid export format."
        )

    # -------------------------------------------------
    # READ DATABASE
    # -------------------------------------------------

    try:

        connection = sqlite3.connect(
            DATABASE_NAME
        )

        query = """
            SELECT
                id,
                age,
                gender,
                year_of_study,
                attendance_percentage,
                average_grade,
                failed_subjects,
                assignments_completed_percentage,
                engagement_score,
                study_hours_per_day,
                online_learning_hours_per_week,
                family_income_inr,
                scholarship,
                part_time_job,
                financial_stress_score,
                distance_from_college_km,
                internet_access,
                parental_support,
                extracurricular_activities,
                disciplinary_issues,
                dropout_probability,
                risk_level,
                prediction,
                created_at
            FROM predictions
        """

        params = ()

        if risk != "ALL":

            query += """
                WHERE UPPER(risk_level) = ?
            """

            params = (risk,)

        query += """
            ORDER BY id DESC
        """

        df = pd.read_sql_query(
            query,
            connection,
            params=params
        )

        connection.close()

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=f"Could not read prediction history: {str(e)}"
        )


    # -------------------------------------------------
    # NO RECORDS
    # -------------------------------------------------

    if df.empty:

        raise HTTPException(
            status_code=404,
            detail=f"No prediction records found for {risk} risk."
        )


    # -------------------------------------------------
    # COLUMN NAMES FOR EXPORT
    # -------------------------------------------------

    export_columns = {
        "id": "Prediction ID",
        "age": "Age",
        "gender": "Gender",
        "year_of_study": "Year of Study",
        "attendance_percentage": "Attendance (%)",
        "average_grade": "Average Grade",
        "failed_subjects": "Failed Subjects",
        "assignments_completed_percentage": "Assignments Completed (%)",
        "engagement_score": "Engagement Score",
        "study_hours_per_day": "Study Hours / Day",
        "online_learning_hours_per_week": "Online Learning Hours / Week",
        "family_income_inr": "Family Income (INR)",
        "scholarship": "Scholarship",
        "part_time_job": "Part Time Job",
        "financial_stress_score": "Financial Stress Score",
        "distance_from_college_km": "Distance From College (KM)",
        "internet_access": "Internet Access",
        "parental_support": "Parental Support",
        "extracurricular_activities": "Extracurricular Activities",
        "disciplinary_issues": "Disciplinary Issues",
        "dropout_probability": "Dropout Probability (%)",
        "risk_level": "Risk Level",
        "prediction": "Prediction",
        "created_at": "Created At"
    }

    df = df.rename(
        columns=export_columns
    )


    # -------------------------------------------------
    # FILE NAME
    # -------------------------------------------------

    risk_name = risk.lower()

    filename = (
        f"{risk_name}_predictions."
        f"{format}"
    )


    # =================================================
    # CSV EXPORT
    # =================================================

    if format == "csv":

        output = io.StringIO()

        df.to_csv(
            output,
            index=False
        )

        output.seek(0)

        return StreamingResponse(
            iter([output.getvalue()]),
            media_type="text/csv",
            headers={
                "Content-Disposition":
                    f'attachment; filename="{filename}"'
            }
        )


    # =================================================
    # EXCEL EXPORT
    # =================================================

    if format == "xlsx":

        output = io.BytesIO()

        workbook = Workbook()

        worksheet = workbook.active

        worksheet.title = "Prediction History"

        # Title
        worksheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=len(df.columns)
        )

        title_cell = worksheet.cell(
            row=1,
            column=1
        )

        title_cell.value = (
            f"StudentGuard AI - "
            f"{risk_name.upper()} Risk Predictions"
        )

        title_cell.font = Font(
            bold=True,
            size=14
        )

        title_cell.alignment = Alignment(
            horizontal="center"
        )

        # Generated time
        worksheet.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=len(df.columns)
        )

        generated_cell = worksheet.cell(
            row=2,
            column=1
        )

        generated_cell.value = (
            "Generated: "
            + datetime.now().strftime(
                "%Y-%m-%d %H:%M:%S"
            )
        )

        generated_cell.alignment = Alignment(
            horizontal="center"
        )

        # Header row
        header_row = 4

        for col_index, column_name in enumerate(
            df.columns,
            start=1
        ):

            cell = worksheet.cell(
                row=header_row,
                column=col_index
            )

            cell.value = column_name

            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="5C1D2B"
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # Data
        for row_index, row in enumerate(
            df.itertuples(index=False),
            start=header_row + 1
        ):

            for col_index, value in enumerate(
                row,
                start=1
            ):

                cell = worksheet.cell(
                    row=row_index,
                    column=col_index
                )

                if pd.isna(value):

                    cell.value = ""

                else:

                    cell.value = value

        # Column widths
        for column_cells in worksheet.columns:

            max_length = 0

            column_letter = (
                column_cells[0].column_letter
            )

            for cell in column_cells:

                try:

                    value_length = len(
                        str(cell.value)
                    )

                    max_length = max(
                        max_length,
                        value_length
                    )

                except Exception:
                    pass

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 2,
                35
            )

        worksheet.freeze_panes = "A5"

        workbook.save(output)

        output.seek(0)

        return StreamingResponse(
            output,
            media_type=(
                "application/"
                "vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition":
                    f'attachment; filename="{filename}"'
            }
        )


    # =================================================
    # PDF EXPORT
    # =================================================

    if format == "pdf":

        output = io.BytesIO()

        document = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            rightMargin=20,
            leftMargin=20,
            topMargin=20,
            bottomMargin=20
        )

        styles = getSampleStyleSheet()

        title_style = styles["Title"]

        title_style.alignment = TA_CENTER

        normal_style = styles["Normal"]

        elements = []

        # Title
        elements.append(
            Paragraph(
                "StudentGuard AI",
                title_style
            )
        )

        elements.append(
            Paragraph(
                "Student Dropout Prediction History",
                styles["Heading2"]
            )
        )

        elements.append(
            Paragraph(
                f"Risk Filter: {risk_name.upper()}",
                normal_style
            )
        )

        elements.append(
            Paragraph(
                "Generated: "
                + datetime.now().strftime(
                    "%Y-%m-%d %H:%M:%S"
                ),
                normal_style
            )
        )

        elements.append(
            Paragraph(
                f"Total Records: {len(df)}",
                normal_style
            )
        )

        elements.append(
            Spacer(
                1,
                12
            )
        )

        # -------------------------------------------------
        # PDF TABLE
        # -------------------------------------------------

        pdf_columns = [
            "Prediction ID",
            "Age",
            "Year of Study",
            "Attendance (%)",
            "Average Grade",
            "Failed Subjects",
            "Dropout Probability (%)",
            "Risk Level",
            "Prediction",
            "Created At"
        ]

        pdf_df = df[
            pdf_columns
        ].copy()

        table_data = [
            pdf_columns
        ]

        for _, row in pdf_df.iterrows():

            table_data.append(
                [
                    str(row[column])
                    if not pd.isna(row[column])
                    else ""
                    for column in pdf_columns
                ]
            )

        table = Table(
            table_data,
            repeatRows=1,
            colWidths=[
                45,
                30,
                55,
                60,
                55,
                55,
                70,
                50,
                145,
                90
            ]
        )

        table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.HexColor(
                            "#5C1D2B"
                        )
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold"
                    ),
                    (
                        "FONTSIZE",
                        (0, 0),
                        (-1, -1),
                        6
                    ),
                    (
                        "ALIGN",
                        (0, 0),
                        (-1, -1),
                        "CENTER"
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "MIDDLE"
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.4,
                        colors.grey
                    ),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [
                            colors.white,
                            colors.HexColor(
                                "#F5F3EF"
                            )
                        ]
                    ),
                    (
                        "TOPPADDING",
                        (0, 0),
                        (-1, -1),
                        4
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, -1),
                        4
                    )
                ]
            )
        )

        elements.append(table)

        document.build(elements)

        output.seek(0)

        return StreamingResponse(
            output,
            media_type="application/pdf",
            headers={
                "Content-Disposition":
                    f'attachment; filename="{filename}"'
            }
        )